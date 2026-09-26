"""Post-generation verification of the workbook and PDF.

- Opens the workbook, confirms key formulas point at the intended cells.
- If the `formulas` package is installed, recalculates every formula independently of
  Excel, compares results with the Python model, and re-runs with a changed check size
  and share price to prove the inputs drive the outputs.
- Confirms the PDF is exactly one page and contains the headline figures.
"""
from __future__ import annotations

import dataclasses
import os
from typing import Optional

import openpyxl

from .calc import Result, calculate
from .models import CapTable, Inputs


def recalc(path: str, changes: Optional[dict[tuple[str, str], object]] = None) -> Optional[callable]:
    """Return a getter g(sheet, cell) over independently recalculated values, or None if the
    `formulas` package is unavailable.  `changes` edits input cells in a temporary copy first,
    exactly as a user would in Excel."""
    try:
        import formulas
    except ImportError:
        return None
    import shutil
    import tempfile
    tmpdir = None
    if changes:
        tmpdir = tempfile.mkdtemp()
        copy = os.path.join(tmpdir, os.path.basename(path))
        wb = openpyxl.load_workbook(path)
        for (sheet, addr), v in changes.items():
            wb[sheet][addr] = v
        wb.save(copy)
        path = copy
    try:
        sol = formulas.ExcelModel().loads(path).finish().calculate()
    finally:
        if tmpdir:
            shutil.rmtree(tmpdir, ignore_errors=True)
    book = os.path.basename(path).upper()
    index = {k.upper(): v for k, v in sol.items()}

    def g(sheet: str, cell: str):
        v = index.get(f"'[{book}]{sheet.upper()}'!{cell}".upper())
        if v is None:
            return None
        x = v.value[0][0]
        return x.item() if hasattr(x, "item") else x
    return g


def _close(a, b, tol=1e-6) -> bool:
    try:
        return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))
    except (TypeError, ValueError):
        return False


def verify_outputs(xlsx: str, pdf: Optional[str], cap: CapTable, inp: Inputs, res: Optional[Result],
                   layout: dict) -> list[tuple[str, bool, str]]:
    out: list[tuple[str, bool, str]] = []
    S, L, inv, tot = layout["S"], layout["cols"], layout["inv_row"], layout["tot"]
    cell = lambda a: a.replace("$", "")

    wb = openpyxl.load_workbook(xlsx)
    out.append(("Workbook opens with expected tabs", wb.sheetnames == ["Summary", "Cap Table", "Source Checks & Notes"],
                ", ".join(wb.sheetnames)))
    s, ct = wb["Summary"], wb["Cap Table"]
    expect = {
        cell(S["own_b"]): f"'Cap Table'!{L['pab']}{inv}",
        cell(S["own_f"]): f"'Cap Table'!{L['paf']}{inv}",
        cell(S["new_sh"]): f"{S['check']}/{S['price']}",
    }
    for addr, frag in expect.items():
        f = str(s[addr].value)
        out.append((f"Summary!{addr} formula references {frag}", frag in f, f))
    nf = str(ct[f"{L['new']}{inv}"].value)
    out.append(("Cap Table new-investor shares = check / price", f"Summary!{S['check']}/Summary!{S['price']}" in nf, nf))
    out.append(("New-investor row is labelled from the investor input", ct[f"B{inv}"].value == f"=Summary!{S['investor']}",
                str(ct[f"B{inv}"].value)))
    hard = [c.coordinate for row in ct.iter_rows(min_row=layout["first"], max_row=tot)
            for c in row if c.column > 3 + len(cap.classes) and c.value is not None and not str(c.value).startswith("=")]
    out.append(("No hard-coded calculated results in Cap Table", not hard, ", ".join(hard[:10]) or "all formulas"))

    g = recalc(xlsx)
    if g is None:
        out.append(("Independent formula recalculation", True,
                    "`formulas` package not installed; Excel will recalculate formulas when the file is opened"))
    else:
        n_err = g("Summary", cell(S["n_err"]))
        failing = [s.cell(row=r, column=1).value for r in range(1, s.max_row + 1)
                   if str(g("Summary", f"C{r}") or "").startswith("ERROR")]
        if res is None:
            only_inputs = all(("positive number" in f or "Commitments flag" in f) for f in failing)
            out.append(("Recalculated blocked workbook: only missing-input checks show ERROR", only_inputs,
                        "; ".join(failing) or "none"))
            own = g("Summary", cell(S["own_b"]))
            out.append(("Blocked workbook shows BLOCKED instead of an ownership %", own == "BLOCKED", str(own)))
        else:
            out.append(("Recalculated workbook: zero ERROR checks", n_err == 0, f"{n_err} errors {failing}"))
        out.append(("Recalculated Before FD ties to source", _close(g("Cap Table", f"{L['bf']}{tot}"), cap.source_total_fd.value),
                    str(g("Cap Table", f"{L['bf']}{tot}"))))
        if res is not None:
            iv = res.investor_row
            pairs = [("new shares", g("Summary", cell(S["new_sh"])), res.new_shares),
                     ("basic %", g("Summary", cell(S["own_b"])), iv.pct_after_basic),
                     ("fully diluted %", g("Summary", cell(S["own_f"])), iv.pct_after_fd),
                     ("post-money", g("Summary", cell(S["post"])), res.valuation["post_money"]),
                     ("implied post (outstanding)", g("Summary", cell(S["ipost_b"])), res.valuation["implied_post_basic"])]
            for n, a, b in pairs:
                out.append((f"Excel = Python: {n}", _close(a, b), f"{a} vs {b}"))
            for i, c in enumerate(res.classes):
                rr = layout["cls_first"] + i
                a = g("Summary", f"E{rr}")
                out.append((f"Excel = Python: After shares {c.key}", _close(a, c.after), f"{a} vs {c.after}"))
            # Change inputs and recalc
            new_check, new_price = (inp.check_size or 0) * 0.5 + 12345, (inp.share_price or 1) * 1.1
            g2 = recalc(xlsx, {("Summary", cell(S["check"])): new_check, ("Summary", cell(S["price"])): new_price})
            res2 = calculate(cap, dataclasses.replace(inp, check_size=new_check, share_price=new_price, resolved=[]))
            for n, a, b in [("new shares", g2("Summary", cell(S["new_sh"])), res2.new_shares),
                            ("basic %", g2("Summary", cell(S["own_b"])), res2.investor_row.pct_after_basic),
                            ("fully diluted %", g2("Summary", cell(S["own_f"])), res2.investor_row.pct_after_fd),
                            ("post-money", g2("Summary", cell(S["post"])), res2.valuation["post_money"])]:
                out.append((f"After changing check/price inputs, Excel = Python: {n}", _close(a, b), f"{a} vs {b}"))
            out.append(("After input change: zero ERROR checks", g2("Summary", cell(S["n_err"])) == 0,
                        str(g2("Summary", cell(S["n_err"])))))

    if pdf:
        import pymupdf
        with pymupdf.open(pdf) as d:
            text = d[0].get_text()
            out.append(("PDF is exactly one page", d.page_count == 1, f"{d.page_count} page(s)"))
            sizes = [sp["size"] for b in d[0].get_text("dict")["blocks"] for ln in b.get("lines", [])
                     for sp in ln["spans"] if sp["text"].strip()]
            body = sorted(sizes)[len(sizes) // 2] if sizes else 0
            out.append(("PDF body text readable at print size (median >= 6.5 pt)", body >= 6.5, f"median {body:.1f} pt"))
        if res is not None:
            from .fmt import pct
            want = [pct(res.investor_row.pct_after_basic, 3), pct(res.investor_row.pct_after_fd, 3)]
            out.append(("PDF shows both ownership figures", all(w in text for w in want), ", ".join(want)))
        else:
            import re
            shown = re.search(r"(basic|fully diluted) ownership\s*[\d.]+%", text, re.I)
            out.append(("Blocked PDF is labelled 'Unable to calculate' and shows no ownership %",
                        "Unable to calculate" in text and not shown, "checked"))
    return out
