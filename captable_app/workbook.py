"""Excel deliverable: Summary, Cap Table and Source Checks & Notes, driven by live formulas.

Only source figures (share counts, recorded cash, source totals) and user inputs are
constants.  Every derived number - new shares, ownership, class totals, valuations,
reconciliation checks - is an Excel formula that recalculates when an input changes.
"""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .calc import NEW_INVESTOR, UNCOUNTED
from .models import CapTable, DeckTerms, Inputs
from .validate import Review

FONT = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFE699")      # highlighted editable inputs
SOURCE_FILL = PatternFill("solid", fgColor="F2F2F2")     # constants copied from source
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
INV_FILL = PatternFill("solid", fgColor="FCE4D6")        # new-investor row
ADJ_FILL = PatternFill("solid", fgColor="DDEBF7")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
AMBER_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
FMT_SH = '#,##0.00;[Red]-#,##0.00'
FMT_USD = '"$"#,##0.00;[Red]-"$"#,##0.00'
FMT_USD4 = '"$"#,##0.0000'
FMT_PCT = '0.000%'
FMT_PP = '+0.000" pp";[Red]-0.000" pp";0.000" pp"'

ROW_TYPE_LABEL = {
    "holder": "Shareholder",
    "commitment": "Seed-3 commitment (already in source pro forma)",
    "placeholder": "Unsold allocation plug (source row)",
    "aggregate": "Aggregate row (holders not named in source)",
    "pool": "Unissued option pool",
    UNCOUNTED: "Adjustment: Seed-3 commitments not in source",
    NEW_INVESTOR: "NEW INVESTOR (this check)",
}


def _f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def _hdr(ws, row, headers, col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill, c.border = _f(True), HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_workbook(path: str, cap: CapTable, deck: DeckTerms, inp: Inputs, review: Review,
                   ai=None, ai_error: str | None = None) -> dict:
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    ct = wb.create_sheet("Cap Table")
    notes = wb.create_sheet("Source Checks & Notes")
    for ws in (s, ct, notes):
        ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------------ Summary: inputs
    S = {}   # named Summary cell addresses
    s["A1"] = f"{cap.company} - Series Seed-3 new-investor ownership"
    s["A1"].font = _f(True, 14)
    s["A2"] = (f"Sources: {cap.file} (cap table as of {cap.as_of}); {deck.file}. Generated {date.today():%Y-%m-%d}. "
               "Yellow cells are editable inputs; grey cells are figures copied from the source.")
    s["A2"].font = _f(italic=True, size=9, color="595959")
    _hdr(s, 5, ["Input", "Value", "Source", "Effect of changing it"])
    in_rows = [
        ("investor", "Investor name", inp.investor, None, "Recalculates (label only)"),
        ("check", "Check size (USD)", inp.check_size, FMT_USD, "Recalculates shares, ownership, valuation, checks"),
        ("price", "Seed-3 share price (USD/share)", inp.share_price, FMT_USD4, "Recalculates shares, ownership, valuation, checks"),
        ("cib", "Existing Seed-3 commitments already in Before? (yes/no)",
         None if inp.commitments_in_before is None else ("yes" if inp.commitments_in_before else "no"), None,
         "Recalculates: 'no' adds the amount below as a separate aggregate adjustment"),
        ("unc", "Seed-3 commitments NOT in Before (USD; used only when 'no')", inp.uncounted_commitments, FMT_USD,
         "Recalculates"),
        ("ph", "Unsold allocation plug row: release or keep", inp.placeholder_mode, None,
         "Recalculates: 'release' removes the plug row in After; 'keep' leaves it (double-count risk)"),
        ("pre", "Pre-money valuation (USD)", inp.pre_money, FMT_USD, "Recalculates post-money"),
    ]
    src_by_name = {r.name: r.source for r in inp.resolved}
    name_map = {"investor": "Investor", "check": "Check size", "price": "Share price",
                "cib": "Commitments already in Before", "unc": "Commitments not in Before (USD)",
                "ph": "Unsold allocation plug row", "pre": "Pre-money valuation"}
    r = 6
    for key, label, val, fmt, effect in in_rows:
        s.cell(row=r, column=1, value=label).font = _f()
        c = s.cell(row=r, column=2, value=val)
        c.fill, c.border, c.font = INPUT_FILL, BOX, _f(True)
        if fmt:
            c.number_format = fmt
        src = src_by_name.get(name_map[key]) or ("REQUIRED - not established by the sources" if val is None else "")
        s.cell(row=r, column=3, value=src).font = _f(size=9)
        s.cell(row=r, column=4, value=effect).font = _f(size=9, color="595959")
        S[key] = f"$B${r}"
        r += 1
    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"release,keep"', allow_blank=False)
    s.add_data_validation(dv)
    s.add_data_validation(dv2)
    dv.add(S["cib"].replace("$", ""))
    dv2.add(S["ph"].replace("$", ""))
    s.cell(row=r, column=1, value="All required inputs present and valid?").font = _f(True)
    S["ok"] = f"$B${r}"
    s.cell(row=r, column=2, value=(
        f"=AND(ISNUMBER({S['check']}),{S['check']}>0,ISNUMBER({S['price']}),{S['price']}>0,"
        f"OR({S['cib']}=\"yes\",AND({S['cib']}=\"no\",ISNUMBER({S['unc']}))),"
        f"OR({S['ph']}=\"release\",{S['ph']}=\"keep\"),ISNUMBER({S['pre']}))")).font = _f(True)
    s.cell(row=r, column=3, value="FALSE = blocked: results show BLOCKED until every required input is set").font = _f(size=9)
    s["A3"] = f'=IF({S["ok"]},"Status: calculated from the inputs below","Status: BLOCKED - required inputs missing (see Source Checks & Notes)")'
    s["A3"].font = _f(True, 11, "C00000")
    r += 2

    # Source constants (grey)
    s.cell(row=r, column=1, value="Figures copied from the source (changing them requires rerunning source extraction)").font = _f(True, 11)
    r += 1
    _hdr(s, r, ["Figure", "Value", "Source location", ""])
    r += 1
    rc = cap.round_class
    open_deck = deck.open_allocation[0] if deck.open_allocation else None
    consts = [
        ("cash_exist", "Seed-3 cash recorded for shares already issued", cap.round_cash_existing, FMT_USD),
        ("cash_commit", "Seed-3 cash for commitment rows already in Before", cap.round_cash_commitments, FMT_USD),
        ("cash_open", "Unsold open allocation per cap table (plug row cash)", cap.round_cash_open, FMT_USD),
        ("src_basic", "Source total shares outstanding", cap.source_total_basic, FMT_SH),
        ("src_fd", "Source total fully diluted shares", cap.source_total_fd, FMT_SH),
        ("src_fd2", "Source total fully diluted shares (Summary tab)", cap.source_total_fd_alt, FMT_SH),
        ("src_price", "Cap-table Seed-3 original issue price", cap.issue_price, FMT_USD4),
        ("deck_open", "Deck open allocation ('Open to New Investors')", open_deck, FMT_USD),
    ]
    for key, label, fig, fmt in consts:
        s.cell(row=r, column=1, value=label).font = _f()
        c = s.cell(row=r, column=2, value=fig.value if fig else 0)
        c.fill, c.border, c.number_format = SOURCE_FILL, BOX, fmt
        s.cell(row=r, column=3, value=str(fig.ref) if fig else "Not found in source (0 used)").font = _f(size=9)
        S[key] = f"$B${r}"
        r += 1
    r += 1

    # ------------------------------------------------------------------ Cap Table sheet
    classes = cap.classes
    ct["A1"] = f"{cap.company} - Cap Table: Before (source pro forma) -> adjustments -> After"
    ct["A1"].font = _f(True, 13)
    ct["A2"] = ("Before = source rows exactly as reported (preferred shown 1:1 as converted). Basic = common + preferred "
                "issued and outstanding; fully diluted adds warrants, options/RSUs and the unissued pool. Adjustments "
                "are Seed-3 shares, shown in separate columns. pp = percentage points.")
    ct["A2"].font = _f(italic=True, size=9, color="595959")
    cols = ["Source row", "Shareholder / line", "Row type"] + [c.name for c in classes]
    ccol = {c.key: get_column_letter(4 + i) for i, c in enumerate(classes)}
    n = 4 + len(classes)
    L = {k: get_column_letter(n + i) for i, k in enumerate(
        ["bb", "bf", "pbb", "pbf", "adjc", "adjp", "new", "ab", "af", "pab", "paf", "db", "df"])}
    cols += ["Before: basic shares", "Before: fully diluted shares", "Before: basic %", "Before: fully diluted %",
             "Adj: Seed-3 commitments not in source", "Adj: release unsold plug row (Seed-3)",
             "New investor Seed-3 shares", "After: basic shares", "After: fully diluted shares", "After: basic %",
             "After: fully diluted %", "Change basic (pp)", "Change fully diluted (pp)"]
    _hdr(ct, 4, cols)
    ct.row_dimensions[4].height = 60
    first = 5
    out_rows = [(row.name, row.row_type, f"{cap.extras['detail_sheet']} row {row.source_row}", row.shares, row)
                for row in cap.rows]
    out_rows.append(("Seed-3 commitments not in source cap table (aggregate; holders not named)", UNCOUNTED,
                     "User input", {}, None))
    out_rows.append((None, NEW_INVESTOR, "User input", {}, None))
    last = first + len(out_rows) - 1
    tot = last + 1
    outstanding_keys = [c.key for c in classes if c.is_outstanding]
    ok = f"Summary!{S['ok']}"
    ph_rows, commit_rows, named_rows = [], [], []
    inv_row = None
    for i, (name, rtype, ref, shares, _src) in enumerate(out_rows):
        rr = first + i
        ct.cell(row=rr, column=1, value=ref)
        ct.cell(row=rr, column=2, value=name if name is not None else f"=Summary!{S['investor']}")
        ct.cell(row=rr, column=3, value=ROW_TYPE_LABEL.get(rtype, rtype))
        for c in classes:
            v = shares.get(c.key, 0.0)
            cell = ct.cell(row=rr, column=4 + classes.index(c), value=v if v else None)
            cell.number_format = FMT_SH
        basic = "+".join(f"N({ccol[k]}{rr})" for k in outstanding_keys)
        ct[f"{L['bb']}{rr}"] = f"={basic}"
        ct[f"{L['bf']}{rr}"] = f"=SUM({ccol[classes[0].key]}{rr}:{ccol[classes[-1].key]}{rr})"
        ct[f"{L['pbb']}{rr}"] = f"={L['bb']}{rr}/{L['bb']}${tot}"
        ct[f"{L['pbf']}{rr}"] = f"={L['bf']}{rr}/{L['bf']}${tot}"
        if rtype == UNCOUNTED:
            ct[f"{L['adjc']}{rr}"] = (f"=IF(AND({ok},Summary!{S['cib']}=\"no\"),"
                                      f"Summary!{S['unc']}/Summary!{S['price']},0)")
        if rtype == "placeholder":
            ct[f"{L['adjp']}{rr}"] = f"=IF(Summary!{S['ph']}=\"release\",-{ccol[rc]}{rr},0)"
            ph_rows.append(rr)
        if rtype == NEW_INVESTOR:
            ct[f"{L['new']}{rr}"] = f"=IF({ok},Summary!{S['check']}/Summary!{S['price']},0)"
            inv_row = rr
        if rtype == "commitment":
            commit_rows.append(rr)
        if rtype in ("holder", "aggregate"):
            named_rows.append(rr)
        adj = f"N({L['adjc']}{rr})+N({L['adjp']}{rr})+N({L['new']}{rr})"
        ct[f"{L['ab']}{rr}"] = f"={L['bb']}{rr}+{adj}"
        ct[f"{L['af']}{rr}"] = f"={L['bf']}{rr}+{adj}"
        ct[f"{L['pab']}{rr}"] = f"=IF({ok},{L['ab']}{rr}/{L['ab']}${tot},\"-\")"
        ct[f"{L['paf']}{rr}"] = f"=IF({ok},{L['af']}{rr}/{L['af']}${tot},\"-\")"
        ct[f"{L['db']}{rr}"] = f"=IF({ok},({L['pab']}{rr}-{L['pbb']}{rr})*100,\"-\")"
        ct[f"{L['df']}{rr}"] = f"=IF({ok},({L['paf']}{rr}-{L['pbf']}{rr})*100,\"-\")"
        fill = INV_FILL if rtype == NEW_INVESTOR else ADJ_FILL if rtype in (UNCOUNTED, "placeholder", "commitment") else None
        for ci in range(1, len(cols) + 1):
            cell = ct.cell(row=rr, column=ci)
            cell.border = BOX
            cell.font = _f(bold=rtype == NEW_INVESTOR, size=9)
            if fill:
                cell.fill = fill
    for key in ("bb", "bf", "adjc", "adjp", "new", "ab", "af"):
        for rr in range(first, tot):
            ct[f"{L[key]}{rr}"].number_format = FMT_SH
    for key in ("pbb", "pbf", "pab", "paf"):
        for rr in range(first, tot):
            ct[f"{L[key]}{rr}"].number_format = FMT_PCT
    for key in ("db", "df"):
        for rr in range(first, tot):
            ct[f"{L[key]}{rr}"].number_format = FMT_PP

    # Totals, source totals, differences
    ct.cell(row=tot, column=2, value="Total (independent sum of rows above)")
    for key in [c.key for c in classes]:
        ct[f"{ccol[key]}{tot}"] = f"=SUM({ccol[key]}{first}:{ccol[key]}{last})"
    for key in ("bb", "bf", "pbb", "pbf", "adjc", "adjp", "new", "ab", "af"):
        ct[f"{L[key]}{tot}"] = f"=SUM({L[key]}{first}:{L[key]}{last})"
    ct[f"{L['pab']}{tot}"] = f"=IF({ok},SUM({L['pab']}{first}:{L['pab']}{last}),\"-\")"
    ct[f"{L['paf']}{tot}"] = f"=IF({ok},SUM({L['paf']}{first}:{L['paf']}{last}),\"-\")"
    srow, drow = tot + 1, tot + 2
    ct.cell(row=srow, column=2, value="Source total (as reported)")
    ct.cell(row=drow, column=2, value="Difference: independent sum - source")
    src_refs = []
    for c in classes:
        figs = cap.source_class_totals.get(c.key) or []
        if c.kind == "option" and not figs:
            figs = []
        if figs:
            ct[f"{ccol[c.key]}{srow}"] = figs[0].value
            ct[f"{ccol[c.key]}{srow}"].fill = SOURCE_FILL
            ct[f"{ccol[c.key]}{drow}"] = f"={ccol[c.key]}{tot}-{ccol[c.key]}{srow}"
            src_refs.append(f"{c.key}: {figs[0].ref.location}")
    ct[f"{L['bb']}{srow}"] = f"=Summary!{S['src_basic']}"
    ct[f"{L['bf']}{srow}"] = f"=Summary!{S['src_fd']}"
    ct[f"{L['bb']}{drow}"] = f"={L['bb']}{tot}-{L['bb']}{srow}"
    ct[f"{L['bf']}{drow}"] = f"={L['bf']}{tot}-{L['bf']}{srow}"
    ct.cell(row=srow, column=3, value="; ".join(src_refs)).font = _f(size=8)
    for rr in (tot, srow, drow):
        for ci in range(1, len(cols) + 1):
            cell = ct.cell(row=rr, column=ci)
            cell.font, cell.border = _f(True, 9), BOX
            if ci >= 4 and cell.number_format == "General":
                cell.number_format = FMT_PCT if get_column_letter(ci) in (L["pbb"], L["pbf"], L["pab"], L["paf"]) else FMT_SH
        ct.cell(row=rr, column=1).fill = HEAD_FILL
    ct.column_dimensions["A"].width = 16
    ct.column_dimensions["B"].width = 44
    ct.column_dimensions["C"].width = 30
    for ci in range(4, len(cols) + 1):
        ct.column_dimensions[get_column_letter(ci)].width = 15
    ct.freeze_panes = "D5"

    # ------------------------------------------------------------------ Summary: results
    CT = "'Cap Table'!"
    S3 = ccol[rc]
    blocked = '"BLOCKED"'
    s.cell(row=r, column=1, value="Results").font = _f(True, 12)
    r += 1
    res_rows = [
        ("new_sh", "New investor Seed-3 shares (check / price, unrounded)", f"=IF({S['ok']},{S['check']}/{S['price']},{blocked})", FMT_SH),
        ("own_b", "New investor ownership - basic (issued and outstanding)", f"=IF({S['ok']},{CT}{L['pab']}{inv_row},{blocked})", FMT_PCT),
        ("own_f", "New investor ownership - fully diluted", f"=IF({S['ok']},{CT}{L['paf']}{inv_row},{blocked})", FMT_PCT),
        ("b_tot", "Before: shares outstanding (basic)", f"={CT}{L['bb']}{tot}", FMT_SH),
        ("a_tot", "After: shares outstanding (basic)", f"=IF({S['ok']},{CT}{L['ab']}{tot},{blocked})", FMT_SH),
        ("bf_tot", "Before: fully diluted shares", f"={CT}{L['bf']}{tot}", FMT_SH),
        ("af_tot", "After: fully diluted shares", f"=IF({S['ok']},{CT}{L['af']}{tot},{blocked})", FMT_SH),
    ]
    for key, label, formula, fmt in res_rows:
        s.cell(row=r, column=1, value=label).font = _f(bold=key in ("own_b", "own_f"))
        c = s.cell(row=r, column=2, value=formula)
        c.number_format, c.border, c.font = fmt, BOX, _f(bold=key in ("own_b", "own_f"))
        S[key] = f"$B${r}"
        r += 1
    r += 1

    # By-class table
    s.cell(row=r, column=1, value="Ownership by security class: Before -> After").font = _f(True, 12)
    r += 1
    _hdr(s, r, ["Class", "Before shares", "Before basic %", "Before fully diluted %", "After shares",
                "After basic %", "After fully diluted %", "Change basic (pp)", "Change fully diluted (pp)"])
    s.row_dimensions[r].height = 30
    r += 1
    cls_first = r
    for c in classes:
        s.cell(row=r, column=1, value=c.name)
        s.cell(row=r, column=2, value=f"={CT}{ccol[c.key]}{tot}")
        after = (f"={CT}{ccol[c.key]}{tot}+{CT}{L['adjc']}{tot}+{CT}{L['adjp']}{tot}+{CT}{L['new']}{tot}"
                 if c.key == rc else f"={CT}{ccol[c.key]}{tot}")
        s.cell(row=r, column=5, value=after)
        if c.is_outstanding:
            s.cell(row=r, column=3, value=f"=B{r}/{CT}{L['bb']}{tot}")
            s.cell(row=r, column=6, value=f"=IF({S['ok']},E{r}/{CT}{L['ab']}{tot},\"-\")")
            s.cell(row=r, column=8, value=f"=IF({S['ok']},(F{r}-C{r})*100,\"-\")")
        else:
            for ci in (3, 6, 8):
                s.cell(row=r, column=ci, value="n/a (excluded from basic)")
        s.cell(row=r, column=4, value=f"=B{r}/{CT}{L['bf']}{tot}")
        s.cell(row=r, column=7, value=f"=IF({S['ok']},E{r}/{CT}{L['af']}{tot},\"-\")")
        s.cell(row=r, column=9, value=f"=IF({S['ok']},(G{r}-D{r})*100,\"-\")")
        r += 1
    cls_last = r - 1
    s.cell(row=r, column=1, value="Total")
    for ci, col in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        s.cell(row=r, column=ci, value=f"=SUM({col}{cls_first}:{col}{cls_last})")
    s.cell(row=r, column=6, value=f"=IF({S['ok']},SUM(F{cls_first}:F{cls_last}),\"-\")")
    s.cell(row=r, column=7, value=f"=IF({S['ok']},SUM(G{cls_first}:G{cls_last}),\"-\")")
    cls_tot = r
    r += 1
    s.cell(row=r, column=1, value="  of which Seed-3: new investor (included above)")
    s.cell(row=r, column=5, value=f"={CT}{L['new']}{inv_row}")
    s.cell(row=r, column=6, value=f"=IF({S['ok']},E{r}/{CT}{L['ab']}{tot},\"-\")")
    s.cell(row=r, column=7, value=f"=IF({S['ok']},E{r}/{CT}{L['af']}{tot},\"-\")")
    inv_cls_row = r
    for rr in range(cls_first, r + 1):
        for ci in range(1, 10):
            cell = s.cell(row=rr, column=ci)
            cell.border = BOX
            cell.font = _f(bold=rr in (cls_tot, inv_cls_row), size=9)
            if ci in (2, 5):
                cell.number_format = FMT_SH
            elif ci in (3, 4, 6, 7):
                cell.number_format = FMT_PCT
            elif ci in (8, 9):
                cell.number_format = FMT_PP
        if rr == inv_cls_row:
            for ci in range(1, 10):
                s.cell(row=rr, column=ci).fill = INV_FILL
    r += 2

    # Valuation
    s.cell(row=r, column=1, value="Valuation").font = _f(True, 12)
    r += 1
    commit_sum = "+".join(f"N({CT}{S3}{x})" for x in commit_rows) or "0"
    named_sum = f"SUM({CT}{S3}{first}:{CT}{S3}{last})-({commit_sum})-" + (
        "(" + "+".join(f"N({CT}{S3}{x})" for x in ph_rows) + ")" if ph_rows else "0")
    ph_after = "+".join(f"N({CT}{S3}{x})+N({CT}{L['adjp']}{x})" for x in ph_rows) or "0"
    val_rows = [
        ("pre_v", "Priced pre-money valuation (input)", f"={S['pre']}", FMT_USD, "Input above"),
        ("proc", "Proceeds included in modelled financing",
         f"={S['cash_exist']}+{S['cash_commit']}+IF({S['ph']}=\"keep\",{S['cash_open']},0)"
         f"+IF({S['cib']}=\"no\",N({S['unc']}),0)+N({S['check']})", FMT_USD,
         "Issued Seed-3 cash + commitment rows + (plug cash if kept) + commitments not in Before + this check"),
        ("post", "Post-money valuation = pre-money + proceeds", f"=IF({S['ok']},B{{pre_v}}+B{{proc}},{blocked})", FMT_USD,
         "Stated-basis post-money"),
        ("pre_sh_b", "Pre-round shares outstanding (Before basic less all Seed-3)", f"={CT}{L['bb']}{tot}-{CT}{S3}{tot}", FMT_SH, ""),
        ("pre_sh_f", "Pre-round fully diluted shares (Before FD less all Seed-3)", f"={CT}{L['bf']}{tot}-{CT}{S3}{tot}", FMT_SH, ""),
        ("ipre_b", "Implied priced pre-money: price x pre-round OUTSTANDING shares (not fully diluted)",
         f"=IF({S['ok']},{S['price']}*B{{pre_sh_b}},{blocked})", FMT_USD, ""),
        ("ipre_f", "Implied priced pre-money: price x pre-round FULLY DILUTED shares",
         f"=IF({S['ok']},{S['price']}*B{{pre_sh_f}},{blocked})", FMT_USD, ""),
        ("ipost_b", "Implied priced post-money: price x After OUTSTANDING shares (not fully diluted)",
         f"=IF({S['ok']},{S['price']}*{S['a_tot']},{blocked})", FMT_USD, ""),
        ("ipost_f", "Implied priced post-money: price x After FULLY DILUTED shares",
         f"=IF({S['ok']},{S['price']}*{S['af_tot']},{blocked})", FMT_USD, ""),
        ("pdiff", "Difference: implied post (outstanding basis) - stated-basis post-money",
         f"=IF({S['ok']},B{{ipost_b}}-B{{post}},{blocked})", FMT_USD, "Explained by the bridge below"),
        ("br1", "  Bridge: implied pre-money (outstanding) - stated pre-money",
         f"=IF({S['ok']},B{{ipre_b}}-B{{pre_v}},{blocked})", FMT_USD, ""),
        ("br2", "  Bridge: issued Seed-3 shares x price - recorded cash",
         f"=IF({S['ok']},{S['price']}*({named_sum})-{S['cash_exist']},{blocked})", FMT_USD, ""),
        ("br3", "  Bridge: commitment rows x price - recorded cash",
         f"=IF({S['ok']},{S['price']}*({commit_sum})-{S['cash_commit']},{blocked})", FMT_USD, ""),
        ("br4", "  Bridge: plug row kept in After x price - its cash",
         f"=IF({S['ok']},IF({S['ph']}=\"keep\",{S['price']}*({ph_after})-{S['cash_open']},0),{blocked})", FMT_USD, ""),
        ("br5", "  Bridge: unexplained remainder",
         f"=IF({S['ok']},B{{pdiff}}-B{{br1}}-B{{br2}}-B{{br3}}-B{{br4}},{blocked})", FMT_USD, "Should be 0"),
    ]
    pos = {}
    for i, (key, *_rest) in enumerate(val_rows):
        pos[key] = r + i
    for key, label, formula, fmt, note in val_rows:
        rr = pos[key]
        s.cell(row=rr, column=1, value=label).font = _f()
        c = s.cell(row=rr, column=2, value=formula.format(**pos) if "{" in formula else formula)
        c.number_format, c.border = fmt, BOX
        s.cell(row=rr, column=3, value=note).font = _f(size=9, color="595959")
        S[key] = f"$B${rr}"
    r += len(val_rows) + 1

    # Checks
    s.cell(row=r, column=1, value="Reconciliation and error checks").font = _f(True, 12)
    r += 1
    _hdr(s, r, ["Check", "Value", "Status", "Detail"])
    r += 1
    cls_diff = "+".join(f"ABS(N({CT}{ccol[c.key]}{drow}))" for c in classes)
    af_rng = f"{CT}{L['ab']}{first}:{CT}{L['af']}{last}"
    checks = [
        ("Before outstanding: independent sum - source", f"={CT}{L['bb']}{drow}",
         "=IF(ABS(B{r})<=0.01,\"OK\",\"ERROR: does not tie\")", "Source: " + (str(cap.source_total_basic.ref) if cap.source_total_basic else "n/a")),
        ("Before fully diluted: independent sum - source", f"={CT}{L['bf']}{drow}",
         "=IF(ABS(B{r})<=0.01,\"OK\",\"ERROR: does not tie\")", "Source: " + (str(cap.source_total_fd.ref) if cap.source_total_fd else "n/a")),
        ("Before fully diluted: independent sum - Summary tab total", f"={CT}{L['bf']}{tot}-{S['src_fd2']}",
         "=IF(ABS(B{r})<=0.01,\"OK\",\"ERROR: does not tie\")", "Source: " + (str(cap.source_total_fd_alt.ref) if cap.source_total_fd_alt else "n/a")),
        ("Class totals: sum of absolute differences vs source", f"={cls_diff}",
         "=IF(B{r}<=0.01,\"OK\",\"ERROR: class totals do not tie\")", "Cap Table difference row"),
        ("By-class Before total - row total", f"=B{cls_tot}-{CT}{L['bf']}{tot}",
         "=IF(ABS(B{r})<=0.01,\"OK\",\"ERROR\")", "Class table vs Cap Table"),
        ("By-class After total - row total", f"=E{cls_tot}-{CT}{L['af']}{tot}",
         "=IF(ABS(B{r})<=0.01,\"OK\",\"ERROR\")", "Class table vs Cap Table"),
        ("After basic % sums to 100%", f"=IF({S['ok']},{CT}{L['pab']}{tot}-1,0)",
         "=IF(ABS(B{r})<0.0000001,\"OK\",\"ERROR\")", ""),
        ("After fully diluted % sums to 100%", f"=IF({S['ok']},{CT}{L['paf']}{tot}-1,0)",
         "=IF(ABS(B{r})<0.0000001,\"OK\",\"ERROR\")", ""),
        ("Smallest After share count (negative = impossible allocation)", f"=MIN({af_rng})",
         "=IF(B{r}>=-0.000001,\"OK\",\"ERROR: negative shares\")", ""),
        ("Check size is a positive number", f"={S['check']}",
         f"=IF(AND(ISNUMBER(B{{r}}),N(B{{r}})>0),\"OK\",\"ERROR: enter a positive check size\")", ""),
        ("Share price is a positive number", f"={S['price']}",
         f"=IF(AND(ISNUMBER(B{{r}}),N(B{{r}})>0),\"OK\",\"ERROR: enter a positive share price\")", ""),
        ("Commitments flag is yes/no (with amount when 'no')", f"={S['cib']}",
         f"=IF(OR({S['cib']}=\"yes\",AND({S['cib']}=\"no\",ISNUMBER({S['unc']}))),\"OK\",\"ERROR: set yes, or no plus amount\")", ""),
        ("Check size vs cap-table open allocation", f"=N({S['check']})-{S['cash_open']}",
         "=IF(B{r}<=0.5,\"OK\",\"WARN: exceeds cap-table open allocation\")", "Cap table: " + (str(cap.round_cash_open.ref) if cap.round_cash_open else "n/a")),
        ("Check size vs deck open allocation", f"=N({S['check']})-{S['deck_open']}",
         "=IF(B{r}<=0.5,\"OK\",\"WARN: exceeds deck open allocation\")", "Deck: " + (str(open_deck.ref) if open_deck else "n/a")),
        ("Share price vs cap-table issue price", f"=N({S['price']})-{S['src_price']}",
         "=IF(ABS(B{r})<0.0000001,\"OK\",\"WARN: price differs from cap-table issue price\")", ""),
        ("Fractional new-investor shares", f"=IF({S['ok']},{S['new_sh']}-INT({S['new_sh']}),0)",
         "=IF(B{r}=0,\"OK\",\"FLAG: fractional - confirm issued share count\")", "No fractional-share convention stated in sources"),
        ("Post-money bridge unexplained remainder", f"=IF({S['ok']},{S['br5']},0)",
         "=IF(ABS(B{r})<=0.01,\"OK\",\"ERROR\")", ""),
    ]
    chk_first = r
    for label, val, status, detail in checks:
        s.cell(row=r, column=1, value=label).font = _f(size=9)
        c = s.cell(row=r, column=2, value=val)
        c.border, c.font = BOX, _f(size=9)
        st = s.cell(row=r, column=3, value=status.replace("{r}", str(r)))
        st.border, st.font = BOX, _f(True, 9)
        s.cell(row=r, column=4, value=detail).font = _f(size=8, color="595959")
        r += 1
    chk_last = r - 1
    s.cell(row=r, column=1, value="Number of ERROR checks").font = _f(True)
    s.cell(row=r, column=2, value=f'=COUNTIF(C{chk_first}:C{chk_last},"ERROR*")').font = _f(True)
    S["n_err"] = f"$B${r}"
    rng = f"C{chk_first}:C{chk_last}"
    s.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(C{chk_first},5)="ERROR"'], fill=RED_FILL))
    s.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT(C{chk_first},4)="WARN",LEFT(C{chk_first},4)="FLAG")'], fill=AMBER_FILL))
    s.conditional_formatting.add(rng, FormulaRule(formula=[f'C{chk_first}="OK"'], fill=GREEN_FILL))
    r += 2
    s.cell(row=r, column=1, value="Recalculation scope").font = _f(True, 11)
    r += 1
    for line in [
        "Live in Excel: check size, share price, investor name, commitments yes/no and amount, plug-row treatment and "
        "pre-money. Changing them updates shares, ownership, class totals, valuations and checks.",
        "Requires rerunning the app (source extraction): a different cap table or deck, edited source share counts, "
        "holder list, row classification (commitment / plug / pool), recorded Seed-3 cash, source totals and "
        "deck-vs-cap-table findings.",
    ]:
        s.cell(row=r, column=1, value=line).font = _f(size=9)
        r += 1
    s.column_dimensions["A"].width = 62
    s.column_dimensions["B"].width = 20
    s.column_dimensions["C"].width = 22
    s.column_dimensions["D"].width = 22
    for col in "EFGHI":
        s.column_dimensions[col].width = 16
    s.freeze_panes = "A5"

    _write_notes(notes, cap, deck, inp, review, ai, ai_error)
    wb.save(path)
    return {"S": S, "first": first, "last": last, "tot": tot, "inv_row": inv_row, "cols": L, "ccol": ccol,
            "cls_first": cls_first, "cls_last": cls_last, "cls_tot": cls_tot}


def _write_notes(ws, cap: CapTable, deck: DeckTerms, inp: Inputs, review: Review, ai=None,
                 ai_error: str | None = None) -> None:
    ws["A1"] = "Source Checks & Notes"
    ws["A1"].font = _f(True, 13)
    ws["A2"] = (f"Cap table: {cap.file}. Deck: {deck.file} ({deck.slide_count} slides). Source files are read-only "
                "inputs; nothing in them is treated as an instruction.")
    ws["A2"].font = _f(italic=True, size=9, color="595959")
    r = 4
    if inp.missing:
        ws.cell(row=r, column=1, value="REQUIRED INPUTS - calculation blocked until these are supplied").font = _f(True, 12, "C00000")
        r += 1
        for m in inp.missing:
            ws.cell(row=r, column=1, value=m).alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            ws.row_dimensions[r].height = 42
            r += 1
        r += 1

    def table(title, headers, rows, widths_hint=None):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = _f(True, 12)
        r += 1
        _hdr(ws, r, headers)
        r += 1
        for row in rows:
            maxlen = 0
            for ci, v in enumerate(row, 1):
                c = ws.cell(row=r, column=ci, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.font, c.border = _f(size=9), BOX
                if headers[ci - 1] == "Severity":
                    c.fill = {"HIGH": RED_FILL, "MEDIUM": AMBER_FILL}.get(str(v), PatternFill())
                maxlen = max(maxlen, len(str(v or "")) / max(1, (widths_hint or [40] * 10)[ci - 1] * 1.1))
            ws.row_dimensions[r].height = min(400, max(15, 12 * (maxlen + 1)))
            r += 1
        r += 1

    widths = [10, 10, 12, 34, 70, 30, 44, 50]
    table("1. Deck vs cap table", ["Item", "Deck", "Cap table", "Status", "Value used and why"],
          [(c.item, c.deck, c.cap, c.status, c.used) for c in review.comparisons], [22, 40, 44, 16, 50])
    issues = sorted(review.issues, key=lambda i: i.rank)
    table("2. Findings: mismatches, formula errors, totals, assumptions (most material first)",
          ["ID", "Severity", "Category", "Finding", "Detail", "Source locations", "Arithmetic",
           "Value used / resolution"],
          [(i.code, i.severity, i.category, i.title, i.detail, "; ".join(i.refs), i.arithmetic, i.resolution)
           for i in issues], widths)
    table("3. Resolved inputs", ["Input", "Value", "Source"],
          [(x.name, x.value if not isinstance(x.value, float) else round(x.value, 6), x.source) for x in inp.resolved],
          [22, 20, 90])
    conventions = [
        ("Before", "The source pro forma exactly as reported, including rows the source adds for the round."),
        ("Basic ownership", "Issued and outstanding common plus preferred (1:1 as converted). Excludes unexercised "
                            "options, unvested/unsettled RSUs, warrants and the unissued pool."),
        ("Fully diluted", "Basic plus warrants, options/RSUs outstanding and the unissued pool. Preferred at 1:1 "
                          "(source conversion columns verified 1:1). Convertibles listed without share counts are "
                          "not added (none stated); see CV1."),
        ("Option pool", "No pool top-up is modelled; none is sourced."),
        ("New shares", "Check size / share price, unrounded. No fractional-share convention is stated; the issued "
                       "share count must be confirmed."),
        ("Commitments", "Rows the source adds for the round without stakeholder IDs are commitments already in "
                        "Before and are not added again. Commitments not in Before are added only when the user "
                        "supplies the amount; they appear as one aggregate row (holders not named)."),
        ("Unsold plug row", "Kept in Before as sourced. 'release' removes it in After through a separate adjustment "
                            "column, so the new investor's shares are not stacked on an unsold placeholder."),
        ("Post-money", "Priced pre-money + proceeds in the modelled financing. Implied priced valuations "
                       "(price x shares) are shown on outstanding and fully diluted bases, labelled separately, with "
                       "a bridge explaining the difference."),
        ("Percentage points", "Change in ownership = After % - Before %, in percentage points."),
    ]
    table("4. Calculation conventions", ["Topic", "Convention"], conventions, [22, 110])
    table("5. Deck slides needing manual verification", ["Slide", "Reason"],
          [(f.split(":")[0], f.split(":", 1)[1].strip()) for f in deck.image_flags], [12, 110])
    if ai is not None or ai_error:
        rows = []
        if ai is not None:
            rows += [("Model", f"{ai.model} (served by {ai.served_by or ai.model}); {ai.images_sent} slide image(s) sent; "
                               f"tokens in/out {ai.usage.get('input_tokens', '?')}/{ai.usage.get('output_tokens', '?')}"),
                     ("Summary", ai.summary or "-"),
                     ("Verified findings", f"{len(ai.findings)} (listed in section 2 with IDs AI-nn)"),
                     ("Discarded findings", ("; ".join(ai.discarded) or "none")
                      + " - discarded because the quoted evidence was not found in the extracted sources")]
            rows += [(f.ref.location, f"{f.note.split(':', 1)[1]} = {f.value:,.2f} - \"{f.ref.label}\" "
                                      "(read from an image; can raise a conflict, never resolves an input)")
                     for f in ai.image_figures]
        if ai_error:
            rows.append(("Status", ai_error))
        table("6. Claude AI review", ["Item", "Detail"], rows, [22, 110])
    ws.column_dimensions["A"].width = 22
    for ci, w in zip("BCDEFGH", [40, 44, 34, 70, 30, 44, 50]):
        ws.column_dimensions[ci].width = w
